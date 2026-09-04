"""Mesclagem de planilhas XLSX."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from services.base import BaseMerger, ProgressCb
from services.exceptions import MergeError
from services.registry import extension_of


def _header_row(worksheet) -> list:
    values = []
    for cell in next(worksheet.iter_rows(min_row=1, max_row=1)):
        values.append(cell.value)
    if not any(value is not None and str(value).strip() != "" for value in values):
        raise MergeError(f"A planilha {worksheet.title} não possui cabeçalho.")
    return values


def _header_key(header: list) -> list[str]:
    return ["" if value is None else str(value).strip().casefold() for value in header]


class ExcelMerger(BaseMerger):
    category = "spreadsheet"

    def validate_content(self, paths: list[Path]) -> None:
        for path in paths:
            if extension_of(path) != ".xlsx":
                raise MergeError(
                    "Arquivos XLS não podem ser mesclados diretamente. Salve como XLSX e tente novamente."
                )
            try:
                workbook = load_workbook(path, read_only=True, data_only=False)
            except InvalidFileException as exc:
                raise MergeError(
                    f"O arquivo {path.name} não é uma planilha XLSX válida.",
                    detail=str(exc),
                ) from exc
            except Exception as exc:
                raise MergeError(
                    f"Não foi possível abrir a planilha {path.name}.",
                    detail=str(exc),
                ) from exc
            if not workbook.sheetnames:
                workbook.close()
                raise MergeError(f"A planilha {path.name} não contém abas.")
            workbook.close()

    def merge(self, paths: list[Path], output: Path, progress: ProgressCb) -> Path:
        progress(10, f"Lendo {paths[0].name}")
        first = load_workbook(paths[0])
        result = Workbook()
        default = result.active
        result.remove(default)

        sheet_headers: dict[str, list] = {}
        for source_sheet in first.worksheets:
            target = result.create_sheet(title=source_sheet.title)
            header = _header_row(source_sheet)
            sheet_headers[source_sheet.title] = header
            for row in source_sheet.iter_rows(values_only=True):
                target.append(list(row))
        first.close()

        for index, path in enumerate(paths[1:], start=1):
            percent = 20 + int(70 * (index / max(len(paths) - 1, 1)))
            progress(percent, f"Mesclando {path.name}")
            workbook = load_workbook(path)
            for source_sheet in workbook.worksheets:
                header = _header_row(source_sheet)
                if source_sheet.title not in sheet_headers:
                    target = result.create_sheet(title=source_sheet.title)
                    sheet_headers[source_sheet.title] = header
                    for row in source_sheet.iter_rows(values_only=True):
                        target.append(list(row))
                    continue
                if _header_key(header) != _header_key(sheet_headers[source_sheet.title]):
                    workbook.close()
                    raise MergeError(
                        f"A aba \"{source_sheet.title}\" em {path.name} possui colunas "
                        "incompatíveis com as demais planilhas."
                    )
                target = result[source_sheet.title]
                for row_index, row in enumerate(source_sheet.iter_rows(values_only=True)):
                    if row_index == 0:
                        continue
                    if any(value is not None and str(value).strip() != "" for value in row):
                        target.append(list(row))
            workbook.close()

        progress(95, "Gravando planilha mesclada")
        result.save(output)
        progress(100, "Mesclagem concluída")
        return output
