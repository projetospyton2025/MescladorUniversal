"""Mesclagem de apresentações PowerPoint PPTX."""

from __future__ import annotations

import io
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from services.base import BaseMerger, ProgressCb
from services.exceptions import MergeError
from services.registry import extension_of


def _open_pptx(path: Path) -> Presentation:
    if extension_of(path) != ".pptx":
        raise MergeError(
            "Arquivos PPT não podem ser mesclados diretamente. Salve como PPTX e tente novamente."
        )
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise MergeError(
            f"O arquivo {path.name} não é uma apresentação PowerPoint válida.",
            detail=str(exc),
        ) from exc
    if len(presentation.slides) == 0:
        raise MergeError(f"A apresentação {path.name} não possui slides.")
    return presentation


def _blank_layout(presentation: Presentation):
    for layout in presentation.slide_layouts:
        name = (layout.name or "").casefold()
        if "blank" in name or "em branco" in name:
            return layout
    index = min(6, len(presentation.slide_layouts) - 1)
    return presentation.slide_layouts[max(index, 0)]


def _clear_shapes(slide) -> None:
    tree = slide.shapes._spTree
    removable = []
    for child in list(tree):
        tag = child.tag
        if tag.endswith(("}sp", "}pic", "}grpSp", "}cxnSp", "}graphicFrame")):
            removable.append(child)
    for child in removable:
        tree.remove(child)


def _append_slide(destination: Presentation, source_slide) -> None:
    dest_slide = destination.slides.add_slide(_blank_layout(destination))
    _clear_shapes(dest_slide)
    for shape in source_slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            image = io.BytesIO(shape.image.blob)
            dest_slide.shapes.add_picture(image, shape.left, shape.top, shape.width, shape.height)
            continue
        dest_slide.shapes._spTree.append(deepcopy(shape.element))


class PptMerger(BaseMerger):
    category = "presentation"

    def validate_content(self, paths: list[Path]) -> None:
        for path in paths:
            _open_pptx(path)

    def merge(self, paths: list[Path], output: Path, progress: ProgressCb) -> Path:
        progress(15, f"Lendo {paths[0].name}")
        merged = _open_pptx(paths[0])
        for index, path in enumerate(paths[1:], start=1):
            percent = 20 + int(70 * (index / max(len(paths) - 1, 1)))
            progress(percent, f"Adicionando {path.name}")
            source = _open_pptx(path)
            for slide in source.slides:
                _append_slide(merged, slide)
        progress(95, "Gravando apresentação mesclada")
        merged.save(str(output))
        progress(100, "Mesclagem concluída")
        return output
