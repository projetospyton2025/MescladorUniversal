"""Gera 2 arquivos por extensão e mescla um formato de cada vez no projeto."""

from __future__ import annotations

import csv
import os
import sys
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from PIL import Image
from pypdf import PdfReader, PdfWriter

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from services.exceptions import MergeError
from services.merge_service import MERGERS
from utils.ffmpeg import ffmpeg_available, ffmpeg_cmd, probe, run_command
from utils.file_utils import format_size
from utils.validation import validate_existing_files


def _usar_ffmpeg_local() -> None:
    if ffmpeg_available():
        return
    encontrados = list((BASE / "ferramentas").glob("**/ffmpeg.exe"))
    if not encontrados:
        return
    pasta = encontrados[0].parent
    os.environ["PATH"] = str(pasta) + os.pathsep + os.environ.get("PATH", "")

BASE = Path(__file__).resolve().parent
ENTRADAS = BASE / "entradas"
SAIDAS = BASE / "saidas"
RELATORIO = BASE / "relatorio.txt"

AUDIO = [".aac", ".flac", ".m4a", ".mp3", ".ogg", ".wav", ".wma"]
VIDEO = [".avi", ".mkv", ".mov", ".mp4", ".webm"]
IMAGEM = [".bmp", ".jpeg", ".jpg", ".png", ".tif", ".tiff", ".webp"]
OUTROS = [".pdf", ".csv", ".xls", ".xlsx", ".doc"]

AUDIO_ENCODERS = {
    ".aac": ["-c:a", "aac", "-b:a", "64k", "-f", "adts"],
    ".flac": ["-c:a", "flac"],
    ".m4a": ["-c:a", "aac", "-b:a", "64k"],
    ".mp3": ["-c:a", "libmp3lame", "-q:a", "9"],
    ".ogg": ["-c:a", "libvorbis", "-q:a", "3"],
    ".wav": ["-c:a", "pcm_s16le"],
    ".wma": ["-c:a", "wmav2", "-b:a", "64k"],
}

VIDEO_ENCODERS = {
    ".avi": ["-c:v", "mpeg4", "-q:v", "8", "-c:a", "libmp3lame", "-q:a", "9"],
    ".mkv": ["-c:v", "libx264", "-preset", "ultrafast", "-crf", "28", "-c:a", "aac", "-b:a", "64k"],
    ".mov": ["-c:v", "libx264", "-preset", "ultrafast", "-crf", "28", "-c:a", "aac", "-b:a", "64k"],
    ".mp4": ["-c:v", "libx264", "-preset", "ultrafast", "-crf", "28", "-c:a", "aac", "-b:a", "64k"],
    ".webm": ["-c:v", "libvpx", "-b:v", "200k", "-c:a", "libvorbis", "-q:a", "3"],
}


def _progress(percent: int, message: str) -> None:
    print(f"    [{percent:3d}%] {message}", flush=True)


def _pasta(ext: str) -> Path:
    pasta = ENTRADAS / ext.lstrip(".").lower()
    pasta.mkdir(parents=True, exist_ok=True)
    return pasta


def _pares(ext: str) -> list[Path]:
    pasta = _pasta(ext)
    return [pasta / f"parte1{ext}", pasta / f"parte2{ext}"]


def _rodar_ffmpeg(args: list[str]) -> None:
    ffmpeg = ffmpeg_cmd()
    if not ffmpeg:
        raise MergeError("FFmpeg não encontrado no PATH.")
    completed = run_command([ffmpeg, "-y", "-hide_banner", "-loglevel", "error", *args])
    if completed.returncode != 0:
        raise MergeError(
            f"Falha ao gerar arquivo de teste: {(completed.stderr or '').strip() or 'erro FFmpeg'}",
            detail=completed.stderr,
        )


def gerar_audio(ext: str) -> list[Path]:
    paths = _pares(ext)
    tones = ("440", "660")
    encoder = AUDIO_ENCODERS[ext]
    for path, freq in zip(paths, tones, strict=True):
        print(f"  gerando {path.name}", flush=True)
        _rodar_ffmpeg(
            [
                "-f",
                "lavfi",
                "-i",
                f"sine=frequency={freq}:sample_rate=44100:duration=0.5",
                *encoder,
                str(path),
            ]
        )
    return paths


def gerar_video(ext: str) -> list[Path]:
    paths = _pares(ext)
    cores = ("red", "blue")
    encoder = VIDEO_ENCODERS[ext]
    for path, cor in zip(paths, cores, strict=True):
        print(f"  gerando {path.name}", flush=True)
        _rodar_ffmpeg(
            [
                "-f",
                "lavfi",
                "-i",
                f"color=c={cor}:s=320x240:r=25:d=2",
                "-f",
                "lavfi",
                "-i",
                "sine=frequency=440:sample_rate=44100:duration=2",
                "-shortest",
                *encoder,
                "-pix_fmt",
                "yuv420p",
                str(path),
            ]
        )
    return paths


def gerar_imagem(ext: str) -> list[Path]:
    paths = _pares(ext)
    cores = ((220, 40, 40), (40, 80, 220))
    formato = {".jpg": "JPEG", ".jpeg": "JPEG", ".tif": "TIFF", ".tiff": "TIFF"}.get(ext, ext.lstrip(".").upper())
    for path, cor in zip(paths, cores, strict=True):
        print(f"  gerando {path.name}", flush=True)
        Image.new("RGB", (160, 100), cor).save(path, format=formato)
    return paths


def gerar_pdf() -> list[Path]:
    paths = _pares(".pdf")
    for path, rotulo in zip(paths, ("Primeira página", "Segunda página"), strict=True):
        print(f"  gerando {path.name}", flush=True)
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=400)
        writer.add_metadata({"/Title": rotulo})
        with path.open("wb") as handle:
            writer.write(handle)
    return paths


def gerar_csv() -> list[Path]:
    paths = _pares(".csv")
    linhas = (
        (("Nome", "Valor"), ("Ana", "1")),
        (("Nome", "Valor"), ("Bia", "2")),
    )
    for path, (cabecalho, linha) in zip(paths, linhas, strict=True):
        print(f"  gerando {path.name}", flush=True)
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerow(cabecalho)
            writer.writerow(linha)
    return paths


def gerar_xlsx() -> list[Path]:
    paths = _pares(".xlsx")
    dados = (("Ana", 1), ("Bia", 2))
    for path, linha in zip(paths, dados, strict=True):
        print(f"  gerando {path.name}", flush=True)
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Dados"
        sheet.append(["Nome", "Valor"])
        sheet.append(list(linha))
        workbook.save(path)
    return paths


def gerar_rejeitado(ext: str) -> list[Path]:
    paths = _pares(ext)
    for path, texto in zip(paths, ("parte um", "parte dois"), strict=True):
        print(f"  gerando {path.name} (conteúdo mínimo para o teste de rejeição)", flush=True)
        path.write_bytes(f"arquivo de teste {texto}{ext}\n".encode("utf-8"))
    return paths


def gerar(ext: str) -> list[Path]:
    if ext in AUDIO:
        return gerar_audio(ext)
    if ext in VIDEO:
        return gerar_video(ext)
    if ext in IMAGEM:
        return gerar_imagem(ext)
    if ext == ".pdf":
        return gerar_pdf()
    if ext == ".csv":
        return gerar_csv()
    if ext == ".xlsx":
        return gerar_xlsx()
    if ext in {".xls", ".doc"}:
        return gerar_rejeitado(ext)
    raise ValueError(ext)


def conferir(ext: str, saida: Path) -> str:
    if not saida.exists() or saida.stat().st_size <= 0:
        return "arquivo de saída vazio"
    if ext in IMAGEM or ext == ".pdf":
        paginas = len(PdfReader(str(saida)).pages)
        if paginas < 2:
            return f"PDF com {paginas} página(s); esperado pelo menos 2"
        return f"PDF com {paginas} páginas, {format_size(saida.stat().st_size)}"
    if ext == ".csv":
        texto = saida.read_text(encoding="utf-8-sig")
        if "Ana" not in texto or "Bia" not in texto:
            return "CSV sem as duas linhas de origem"
        return f"CSV com as duas linhas, {format_size(saida.stat().st_size)}"
    if ext == ".xlsx":
        livro = load_workbook(saida, read_only=True, data_only=True)
        valores = [linha[0] for linha in livro.active.iter_rows(values_only=True)]
        livro.close()
        if valores != ["Nome", "Ana", "Bia"]:
            return f"XLSX com linhas inesperadas: {valores}"
        return f"XLSX com cabeçalho + Ana + Bia, {format_size(saida.stat().st_size)}"
    if ext in AUDIO or ext in VIDEO:
        info = probe(saida)
        duracao = float((info.get("format") or {}).get("duration") or 0)
        return f"{saida.suffix.upper().lstrip('.')} {format_size(saida.stat().st_size)}, duração {duracao:.2f}s"
    return format_size(saida.stat().st_size)


def mesclar(ext: str, entradas: list[Path]) -> tuple[bool, str, Path | None]:
    nome = f"mesclado_{ext.lstrip('.').lower()}"
    try:
        detected = validate_existing_files(entradas)
        saida = SAIDAS / f"{nome}{detected['output_ext']}"
        saida.unlink(missing_ok=True)
        merger = MERGERS[detected["category"]]
        merger.validate_content(entradas)
        merger.merge(entradas, saida, _progress)
        detalhe = conferir(ext, saida)
        return True, detalhe, saida
    except MergeError as exc:
        return False, exc.user_message, None


def main() -> int:
    SAIDAS.mkdir(parents=True, exist_ok=True)
    ENTRADAS.mkdir(parents=True, exist_ok=True)
    _usar_ffmpeg_local()

    if "--somente-midia" in sys.argv:
        extensoes = AUDIO + VIDEO
    else:
        extensoes = AUDIO + VIDEO + IMAGEM + OUTROS
    linhas = [
        "Teste de mesclagem por extensão",
        f"Pasta: {BASE}",
        f"FFmpeg: {'ok' if ffmpeg_available() else 'AUSENTE'}",
        "",
    ]
    falhas = 0

    print(f"Saídas em: {SAIDAS}", flush=True)
    print(f"Entradas em: {ENTRADAS}", flush=True)

    for ext in extensoes:
        rotulo = ext.lstrip(".").upper()
        print(f"\n=== {rotulo} ===", flush=True)
        inicio = time.perf_counter()
        try:
            arquivos = gerar(ext)
            for path in arquivos:
                print(f"  entrada: {path} ({format_size(path.stat().st_size)})", flush=True)
            ok, detalhe, saida = mesclar(ext, arquivos)
        except Exception as exc:
            ok, detalhe, saida = False, str(exc), None
        elapsed = time.perf_counter() - inicio
        if ok:
            print(f"  OK  {saida.name} — {detalhe} ({elapsed:.1f}s)", flush=True)
            linhas.append(f"OK   {rotulo:<5}  {saida}  {detalhe}")
        else:
            falhas += 1
            esperado = ext in {".xls", ".doc"}
            marca = "ESP." if esperado else "FALHA"
            if esperado:
                falhas -= 1
            print(f"  {marca} {detalhe} ({elapsed:.1f}s)", flush=True)
            linhas.append(f"{marca:<4} {rotulo:<5}  {detalhe}")

    RELATORIO.write_text("\n".join(linhas) + "\n", encoding="utf-8")
    print(f"\nRelatório: {RELATORIO}", flush=True)
    return 1 if falhas else 0


if __name__ == "__main__":
    raise SystemExit(main())
