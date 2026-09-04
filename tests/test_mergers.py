"""Testes dos processadores de mesclagem e da validação de formatos."""

from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from pypdf import PdfReader, PdfWriter
from PIL import Image
from docx import Document
from openpyxl import Workbook, load_workbook

from services.csv_merger import CsvMerger
from services.doc_merger import DocMerger
from services.excel_merger import ExcelMerger
from services.exceptions import MergeError
from services.image_merger import ImageMerger
from services.json_merger import JsonMerger
from services.pdf_merger import PdfMerger
from services.registry import detect_from_names
from services.text_merger import TextMerger
from utils.ffmpeg import ffmpeg_available


def _noop_progress(_percent: int, _message: str) -> None:
    return None


class DetectTests(unittest.TestCase):
    def test_json_detection(self) -> None:
        info = detect_from_names(["a.json", "b.json"])
        self.assertEqual(info["category"], "json")
        self.assertEqual(info["format_label"], "JSON")

    def test_mp3_detection(self) -> None:
        info = detect_from_names(["a.mp3", "b.mp3", "c.mp3"])
        self.assertEqual(info["category"], "audio")
        self.assertEqual(info["format_label"], "MP3")

    def test_mixed_rejected(self) -> None:
        with self.assertRaises(MergeError) as ctx:
            detect_from_names(["a.mp3", "b.mp4", "c.json"])
        self.assertIn("incompatíveis", ctx.exception.user_message)

    def test_images_become_pdf(self) -> None:
        info = detect_from_names(["a.jpg", "b.png"])
        self.assertEqual(info["output_ext"], ".pdf")

    def test_xlsx_supported(self) -> None:
        info = detect_from_names(["a.xlsx", "b.xlsx"])
        self.assertEqual(info["category"], "spreadsheet")

    def test_txt_supported(self) -> None:
        info = detect_from_names(["a.txt", "b.txt"])
        self.assertEqual(info["category"], "text")

    def test_doc_legacy_rejected(self) -> None:
        with self.assertRaises(MergeError):
            detect_from_names(["a.doc", "b.doc"])


class JsonMergeTests(unittest.TestCase):
    def test_concat_arrays(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "a.json").write_text("[1, 2]", encoding="utf-8")
            (root / "b.json").write_text("[3]", encoding="utf-8")
            output = root / "out.json"
            JsonMerger().merge([root / "a.json", root / "b.json"], output, _noop_progress)
            self.assertEqual(json.loads(output.read_text(encoding="utf-8")), [1, 2, 3])

    def test_merge_objects(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "a.json").write_text('{"nome": "Ana", "itens": [1]}', encoding="utf-8")
            (root / "b.json").write_text('{"cidade": "SP", "itens": [2]}', encoding="utf-8")
            output = root / "out.json"
            JsonMerger().merge([root / "a.json", root / "b.json"], output, _noop_progress)
            data = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual(data["nome"], "Ana")
            self.assertEqual(data["cidade"], "SP")
            self.assertEqual(data["itens"], [1, 2])

    def test_incompatible_structures(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "a.json").write_text("[1]", encoding="utf-8")
            (root / "b.json").write_text('{"a": 1}', encoding="utf-8")
            with self.assertRaises(MergeError):
                JsonMerger().validate_content([root / "a.json", root / "b.json"])


class TextCsvPdfTests(unittest.TestCase):
    def test_txt_concat_order(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "a.txt").write_text("um", encoding="utf-8")
            (root / "b.txt").write_text("dois\n", encoding="utf-8")
            output = root / "out.txt"
            TextMerger().merge([root / "a.txt", root / "b.txt"], output, _noop_progress)
            self.assertEqual(output.read_text(encoding="utf-8"), "um\ndois\n")

    def test_csv_headers_must_match(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "a.csv").write_text("nome,idade\nAna,20\n", encoding="utf-8")
            (root / "b.csv").write_text("nome,cidade\nAna,SP\n", encoding="utf-8")
            with self.assertRaises(MergeError):
                CsvMerger().validate_content([root / "a.csv", root / "b.csv"])

    def test_csv_concat_rows(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            (root / "a.csv").write_text("nome,idade\nAna,20\n", encoding="utf-8")
            (root / "b.csv").write_text("nome,idade\nBia,21\n", encoding="utf-8")
            output = root / "out.csv"
            CsvMerger().merge([root / "a.csv", root / "b.csv"], output, _noop_progress)
            text = output.read_text(encoding="utf-8-sig")
            self.assertIn("Ana,20", text)
            self.assertIn("Bia,21", text)
            self.assertEqual(text.count("nome,idade"), 1)

    def test_pdf_page_count(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for name in ("a.pdf", "b.pdf"):
                writer = PdfWriter()
                writer.add_blank_page(width=72, height=72)
                with (root / name).open("wb") as handle:
                    writer.write(handle)
            output = root / "out.pdf"
            PdfMerger().merge([root / "a.pdf", root / "b.pdf"], output, _noop_progress)
            self.assertEqual(len(PdfReader(str(output)).pages), 2)

    def test_images_to_pdf(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            Image.new("RGB", (20, 20), "red").save(root / "a.jpg")
            Image.new("RGB", (24, 18), "blue").save(root / "b.png")
            output = root / "out.pdf"
            ImageMerger().merge([root / "a.jpg", root / "b.png"], output, _noop_progress)
            self.assertGreaterEqual(len(PdfReader(str(output)).pages), 2)

    def test_xlsx_concat_rows(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for name, rows in (
                ("a.xlsx", [("Nome", "Valor"), ("Ana", 1)]),
                ("b.xlsx", [("Nome", "Valor"), ("Bia", 2)]),
            ):
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "Dados"
                for row in rows:
                    sheet.append(list(row))
                workbook.save(root / name)
            output = root / "out.xlsx"
            ExcelMerger().merge([root / "a.xlsx", root / "b.xlsx"], output, _noop_progress)
            merged = load_workbook(output)
            values = [row[0] for row in merged.active.iter_rows(values_only=True)]
            self.assertEqual(values, ["Nome", "Ana", "Bia"])

    def test_docx_concat(self) -> None:
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for name, text in (("a.docx", "Primeiro"), ("b.docx", "Segundo")):
                document = Document()
                document.add_paragraph(text)
                document.save(root / name)
            output = root / "out.docx"
            DocMerger().merge([root / "a.docx", root / "b.docx"], output, _noop_progress)
            merged = Document(output)
            texts = [paragraph.text for paragraph in merged.paragraphs if paragraph.text]
            self.assertTrue(any("Primeiro" in item for item in texts))
            self.assertTrue(any("Segundo" in item for item in texts))


@unittest.skipUnless(ffmpeg_available(), "FFmpeg não está instalado")
class MediaTests(unittest.TestCase):
    def test_ffmpeg_detected(self) -> None:
        self.assertTrue(ffmpeg_available())

    def test_mp3_concat(self) -> None:
        from services.audio_merger import AudioMerger
        from utils.ffmpeg import ffmpeg_cmd, run_command

        ffmpeg = ffmpeg_cmd()
        self.assertIsNotNone(ffmpeg)
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            for name in ("a.mp3", "b.mp3"):
                completed = run_command(
                    [
                        ffmpeg,
                        "-y",
                        "-f",
                        "lavfi",
                        "-i",
                        "anullsrc=r=44100:cl=stereo",
                        "-t",
                        "0.4",
                        "-c:a",
                        "libmp3lame",
                        "-q:a",
                        "9",
                        str(root / name),
                    ]
                )
                self.assertEqual(completed.returncode, 0)
            output = root / "out.mp3"
            AudioMerger().merge([root / "a.mp3", root / "b.mp3"], output, _noop_progress)
            self.assertTrue(output.exists())
            self.assertGreater(output.stat().st_size, 0)


if __name__ == "__main__":
    unittest.main()
